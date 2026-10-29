import re,shutil,subprocess,tempfile,hashlib,docx,pymupdf,os
from collections import Counter
from pathlib import Path
from docx.table import Table
from pptx import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture
from .vision import caption,prepare


EXTS={".pdf",".docx",".doc",".md",".pptx",".ppt"}
SKIP_DIRS={"venv", "env", "site-packages", "node_modules", "__pycache__",
             "ml_researcher", "System Volume Information"}
CHUNK,OVERLAP=1500,300

def to_wsl_path(raw:str) -> Path:
    """Accept 'E:\\AI_ML\\Notes' or '/mnt/e/AI_ML/Notes'; return a resolved Linux path."""
    m=re.match(r"^([A-Za-z]):[\\/]?(.*)$",raw)
    if m:
        raw=f"/mnt/{m.group(1).lower()}/" + m.group(2).replace("\\","/")
    return Path(raw).expanduser().resolve()

MIN_TEXT=200       # less text than this -> treat as scanned page
VISION_MAX_PAGES=150  
MIN_IMG_FRAC=0.08  # image covering >8% of the page counts as a real figure
MIN_DRAWINGS=50    # vector paths beyond the PDF's baseline -> probably a chart/diagram

def walk(root:Path):
    """Yield every supported file under root, skipping SKIP_DIRS and hidden folders."""
    for dirpath,dirnames,filenames in os.walk(root):
        dirnames[:]=[d for d in dirnames if d not in SKIP_DIRS and not d.startswith(".")]
        for f in filenames:
            if Path(f).suffix.lower() in EXTS and not f.startswith("~$"):  # ~$ = Office lock files
                yield Path(dirpath)/f


def _needs_vision(page,text:str,extra_drawings:int)->bool:
    if len(text)<MIN_TEXT:
        return True
    area=page.rect.width*page.rect.height
    for info in page.get_image_info():
        r=pymupdf.Rect(info["bbox"])
        if r.width*r.height/area>MIN_IMG_FRAC:
            return True
    return extra_drawings>MIN_DRAWINGS

def _read_pdf(p:Path)->str:
    parts=[]
    with pymupdf.open(p) as doc:
        drawings=[len(pg.get_drawings()) for pg in doc] if len(doc)<=VISION_MAX_PAGES else [0]*len(doc)
        base=min(drawings,default=0)     # plainest page = template/background paths
        for n,page in enumerate(doc,1):
            text=page.get_text().strip()
            if len(doc)<=VISION_MAX_PAGES and _needs_vision(page,text,drawings[n-1]-base):
                png=page.get_pixmap(dpi=110).tobytes("png")
                scanned=len(text)<MIN_TEXT
                cap=caption(png,"image/png",ctx=text,ocr=scanned)
                if cap:
                    text=cap if scanned else f"{text}\n\n[Figure description, page {n}]\n{cap}"
            parts.append(text)
    return "\n\n".join(parts)


def _pptx_shapes(shapes,texts:list,imgs:list):
    for sh in shapes:
        if isinstance(sh,GroupShape):
            _pptx_shapes(sh.shapes,texts,imgs)
        elif isinstance(sh,Picture):
            imgs.append((sh.image.blob,sh.image.content_type))
        elif isinstance(sh,GraphicFrame) and sh.has_table:
            for row in sh.table.rows:
                texts.append(" | ".join(c.text for c in row.cells))
        elif isinstance(sh,GraphicFrame) and sh.has_chart:
            for plot in sh.chart.plots:  # native charts: use the real data, no vision needed
                cats=list(plot.categories)
                for s in plot.series:
                    texts.append(f"{s.name}: "+", ".join(f"{c}={v}" for c,v in zip(cats,s.values,strict=False)))  # a series may have fewer values than categories
        elif isinstance(sh,Shape) and sh.has_text_frame:
            texts.append(sh.text_frame.text)

def _read_pptx(p:Path)->str:
    slides=[]
    for slide in Presentation(str(p)).slides:
        texts,imgs=[],[]
        _pptx_shapes(slide.shapes,texts,imgs)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            texts.append(slide.notes_slide.notes_text_frame.text)
        slides.append((texts,imgs))
    counts=Counter(hashlib.sha256(b).hexdigest() for _,imgs in slides for b,_ in imgs)
    parts=[]
    for n,(texts,imgs) in enumerate(slides,1):
        body="\n".join(t for t in texts if t.strip())
        for blob,mime in imgs:
            if counts[hashlib.sha256(blob).hexdigest()]>=3:  # same image on 3+ slides = template logo
                continue
            img=prepare(blob,mime)
            if img and (cap:=caption(*img,ctx=body)):
                body+=f"\n\n[Image on slide {n}]\n{cap}"
        parts.append(f"[Slide {n}]\n{body}")
    return "\n\n".join(parts)

def _read_docx(p:Path)->str:
    d=docx.Document(str(p))
    parts=[]
    for block in d.iter_inner_content():
        if isinstance(block,Table):
            parts.append("\n".join(" | ".join(c.text for c in row.cells) for row in block.rows))
            continue
        if block.text.strip():
            parts.append(block.text)
        for rid in block._p.xpath(".//a:blip/@r:embed"):  # images anchored in this paragraph
            part=d.part.related_parts[rid]
            img=prepare(part.blob,part.content_type)
            if img and (cap:=caption(*img,ctx="\n".join(parts[-3:]))):
                parts.append(f"[Image]\n{cap}")
    return "\n\n".join(parts)

def read_file(p:Path)->str:
    s=p.suffix.lower()
    if s ==".pdf":
        return _read_pdf(p)
    if s ==".docx":
        return _read_docx(p)
    if s == ".pptx" or s ==".ppt":
        if s ==".ppt":
            if not shutil.which("soffice"):
                raise RuntimeError("LibreOffice not installed (sudo apt-get install libreoffice-impress)")
            with tempfile.TemporaryDirectory() as tmp:
                subprocess.run(["soffice","--headless","--convert-to","pptx","--outdir",tmp,str(p)],
                       check=True,capture_output=True,timeout=300)
                return read_file(Path(tmp)/(p.stem+".pptx"))
        return _read_pptx(p)       
    if s ==".doc":
        if not shutil.which("antiword"):
            raise RuntimeError("antiword not installed (sudo apt-get install antiword)")
        return subprocess.run(["antiword",str(p)],capture_output=True,timeout=120,
                              encoding="utf-8",errors="ignore").stdout
    return p.read_text(errors="ignore")

def _split_long(para:str):
    """Split an oversized paragraph into sentence-packed pieces <= CHUNK."""
    buff=""
    for sent in re.split(r"(?<=[.!?])\s+",para):
        if len(sent)>CHUNK: # one giant "sentence" 
            if buff:
                yield buff
                buff=""
            for i in range(0,len(sent),CHUNK):
                yield sent[i:i+CHUNK]
            continue
        if len(buff)+len(sent)+1>CHUNK:
            yield buff
            buff=""
        buff=f"{buff} {sent}".strip()
    if buff:
        yield buff

def chunks(text:str):
    units=[]
    for para in re.split(r"\n\s*\n",text):
        para=para.strip()
        if not para:
            continue
        units.extend([para] if len(para)<CHUNK else _split_long(para))

    buf,size=[],0
    for u in units:
        if buf and size +len(u)>CHUNK:
            piece="\n\n".join(buf)
            if len(piece)>100:
                yield piece
            keep,k=[],0
            for x in reversed(buf):
                if k+len(x)>OVERLAP:
                    break
                keep.insert(0,x)
                k+=len(x)
            buf,size=keep,k
        buf.append(u)
        size += len(u)+2
    piece="\n\n".join(buf)
    if len(piece)>100:
        yield piece

